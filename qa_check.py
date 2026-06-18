# -*- coding: utf-8 -*-
"""Programmatic QA for the deck (renderer-free fallback per nature-paper2ppt self-review)."""
import math, re
from pptx import Presentation
from pptx.util import Emu

PATH = "output/final_presentation_cn.pptx"
SW, SH = 13.333, 7.5
EMU = 914400

prs = Presentation(PATH)

def adv(ch, sz):
    o = ord(ch)
    if o == 0x20: return sz*0.28
    if 0x2E80 <= o <= 0x9FFF or 0xFF00 <= o <= 0xFFEF or o in (0x2014,0x2026,0x201C,0x201D):
        return sz*1.03          # CJK + fullwidth punctuation
    if ch.isdigit(): return sz*0.56
    if ch.isalpha(): return sz*0.56
    return sz*0.45

def para_runs(p):
    return [(r.text, (r.font.size.pt if r.font.size else 18)) for r in p.runs]

issues=[]
def flag(sev, sl, msg): issues.append((sev, sl, msg))

PLACEHOLDER = re.compile(r'lorem|ipsum|\bx{3,}\b|\bTODO\b|\[insert', re.I)

for idx, slide in enumerate(prs.slides, 1):
    for shp in slide.shapes:
        # bounds
        if shp.left is None or shp.top is None: continue
        l=shp.left/EMU; t=shp.top/EMU
        w=(shp.width or 0)/EMU; h=(shp.height or 0)/EMU
        if l < -0.03 or t < -0.03 or l+w > SW+0.05 or t+h > SH+0.05:
            flag("high", idx, f"shape out of bounds: L{l:.2f} T{t:.2f} R{l+w:.2f} B{t+h:.2f} ({shp.shape_type})")
        if not shp.has_text_frame: continue
        tf = shp.text_frame
        ml=(tf.margin_left or 0)/EMU; mr=(tf.margin_right or 0)/EMU
        mt=(tf.margin_top or 0)/EMU; mb=(tf.margin_bottom or 0)/EMU
        usable_w = max(0.1,(w-ml-mr))*72
        usable_h = max(0.1,(h-mt-mb))*72
        wrap = tf.word_wrap is not False
        total_h = 0; maxsingle=0
        for p in tf.paragraphs:
            runs = para_runs(p)
            if not runs:
                total_h += 10; continue
            txt = "".join(r[0] for r in runs)
            if PLACEHOLDER.search(txt):
                flag("high", idx, f"placeholder text: {txt[:30]!r}")
            line_adv = sum(adv(c, sz) for tt,sz in runs for c in tt)
            mx = max(sz for _,sz in runs)
            ls = p.line_spacing if isinstance(p.line_spacing,(int,float)) else 1.0
            lh = mx*1.2*(ls if ls else 1.0)
            sa = p.space_after.pt if p.space_after else 0
            if wrap:
                lines = max(1, math.ceil(line_adv/usable_w - 1e-6))
            else:
                lines = 1
                maxsingle = max(maxsingle, line_adv)
            total_h += lines*lh + sa
        # width overflow (no-wrap single line)
        if not wrap and maxsingle > usable_w*1.02:
            flag("high", idx, f"text wider than box (no-wrap): need {maxsingle/72:.2f}in > {usable_w/72:.2f}in :: {tf.text[:24]!r}")
        # height overflow
        if total_h > usable_h*1.08 and tf.text.strip():
            flag("medium", idx, f"text may overflow height: need~{total_h/72:.2f}in > box {usable_h/72:.2f}in :: {tf.text[:24]!r}")

# structure summary
nslides=len(prs.slides._sldIdLst)
nnotes=sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
print(f"slides={nslides}  slides_with_notes={nnotes}")
hi=[i for i in issues if i[0]=='high']
md=[i for i in issues if i[0]=='medium']
print(f"HIGH={len(hi)}  MEDIUM={len(md)}")
for sev in ('high','medium'):
    for s,sl,m in [i for i in issues if i[0]==sev]:
        print(f"  [{sev:6}] slide {sl}: {m}")
